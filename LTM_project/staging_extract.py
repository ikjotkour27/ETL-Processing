import os
import re
import csv
import hashlib
import mysql.connector
from datetime import datetime
import json
import openpyxl
from dateutil import parser

MANDATORY_COLUMNS = {
    "employee": ["emp_id"],
    "course": ["course_id"],
    "trainer": ["trainer_id"],
    "enrollment": ["enrollment_id", "emp_id", "course_id"],
    "assessment": ["assessment_id", "emp_id", "course_id"]
}

ARCHIVE_DIR = "archive"

DB_CONFIG = {
    "host": "localhost",
    "user": "root",
    "password": "Ikjot123@",
    "database": "ETL_Project",
}

TABLE_CONFIG = {
    "employee":   ("landing_employee",   "emp_id",        "emp_"),
    "course":     ("landing_course",     "course_id",     "course_"),
    "trainer":    ("landing_trainer",    "trainer_id",    "trainer_"),
    "enrollment": ("landing_enrollment", "enrollment_id", "enrollment_"),
    "assessment": ("landing_assessment", "assessment_id", "assessment_"),
}

DATE_PAIR_MAP = {
    "trainer": ("training_start_date", "training_end_date"),
    "course":  ("course_start_date",   "course_end_date"),
}


SCRIPT_NAME = "staging_extract.py"


COLUMN_RENAME_MAP = {
    "employee_id": "emp_id",
    "emp_i_d":     "emp_id",
    "empid":       "emp_id",
    "course_id_fk":"course_id",
    "trainer_id_fk":"trainer_id",
    "enrollment_id_fk": "enrollment_id",
    "assessment_id_fk": "assessment_id",
}


def mandatory_fields_present(row, table_name):
    mandatory = MANDATORY_COLUMNS.get(table_name, [])

    for col in mandatory:
        if row.get(col) is None:
            return False

    return True

def get_latest_load_key(conn):
    cursor = conn.cursor()
    cursor.execute("""
        SELECT load_key
        FROM load_control
        WHERE project_name='LTM'
        ORDER BY load_date DESC
        LIMIT 1
    """)
    row = cursor.fetchone()
    cursor.close()
    return row[0] if row else None


def fetch_landing_data(conn, landing_table, load_key):
    cursor = conn.cursor(dictionary=True)
    cursor.execute(f"""
        SELECT * FROM {landing_table}
        WHERE load_key = %s
    """, (load_key,))
    rows = cursor.fetchall()
    cursor.close()
    return rows


def load_archive_file(prefix):
    files = [f for f in os.listdir(ARCHIVE_DIR) if f.startswith(prefix)]
    if not files:
        return None
    files.sort(key=lambda f: get_date_from_filename(f) or "", reverse=True)
    fpath = os.path.join(ARCHIVE_DIR, files[0])

    if fpath.endswith(".csv"):
        with open(fpath, newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            rows = [standardize_column_casing(row) for row in reader]
        return rows
    elif fpath.endswith(".json"):
        with open(fpath, encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, dict):
            data = [data]
        rows = [standardize_column_casing(row) for row in data]
        return rows
    elif fpath.endswith(".xlsx"):
        wb = openpyxl.load_workbook(fpath, data_only=True)
        sheet = wb.active
        rows_iter = sheet.iter_rows(values_only=True)
        headers = next(rows_iter)
        headers = [str(h) if h is not None else "" for h in headers]
        rows = []
        for values in rows_iter:
            row = dict(zip(headers, values))
            rows.append(standardize_column_casing(row))
        return rows
    return None


def get_active_base_record(conn, table_name, id_col, id_value):
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute(f"""
            SELECT * FROM base_{table_name}
            WHERE `{id_col}` = %s AND curr_ind = 'Y'
        """, (str(id_value),))
        row = cursor.fetchone()
        cursor.close()
        return row
    except Exception:
        cursor.close()
        return None


def create_staging_table(conn, table_name, columns):
    cursor = conn.cursor()
    col_defs = ",\n    ".join([f"`{col}` TEXT" for col in columns])
    cursor.execute(f"""
        CREATE TABLE IF NOT EXISTS staging_{table_name} (
            {table_name}_key VARCHAR(100) PRIMARY KEY,
            {col_defs},
            checksum VARCHAR(32),
            change_flag CHAR(1)
        )
    """)
    conn.commit()
    cursor.close()


def create_audit_table_2(conn):
    cursor = conn.cursor()
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS audit_table_2 (
            audit_id INT AUTO_INCREMENT PRIMARY KEY,
            table_name VARCHAR(100),
            file_name VARCHAR(100),
            ingest_timestamp DATETIME,
            load_key VARCHAR(50),
            status VARCHAR(20),
            file_path VARCHAR(255),
            row_count INT
        )
        """
    )
    conn.commit()
    cursor.close()


def write_to_audit_2(conn, table_name, file_name, ts, load_key, status, file_path, row_count):
    cursor = conn.cursor()
    cursor.execute(
        """
        INSERT INTO audit_table_2 (table_name, file_name, ingest_timestamp, load_key, status, file_path, row_count)
        VALUES (%s, %s, %s, %s, %s, %s, %s)
        """,
        (table_name, file_name, ts, load_key, status, file_path, row_count),
    )
    conn.commit()
    cursor.close()


def find_archive_match(archive_rows, id_col, id_value):
    for row in archive_rows:
        if str(row.get(id_col)) == str(id_value):
            return row
    return None


CHECKSUM_EXCLUDE_COLUMNS = {
    "ingest_timestamp",
    "load_key",
    "staging_key"
}


def compute_checksum(curr_row):
    def normalize(v):
        if v is None:
            return ""
        s = str(v).strip()
        return "" if s.lower() in ("none", "nan", "null") else s

    checksum_cols = [
        k for k in sorted(curr_row.keys())
        if k.lower() not in CHECKSUM_EXCLUDE_COLUMNS
    ]
    row_str = "|".join(normalize(curr_row[k]) for k in checksum_cols)
    return hashlib.md5(row_str.encode("utf-8")).hexdigest()


def make_staging_key(id_value, timestamp):
    ts_str = timestamp.strftime("%Y%m%d%H%M%S")
    return f"{id_value}_{ts_str}"


def get_existing_staging_checksums(conn, table_name, id_col):
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute(f"SELECT `{id_col}`, checksum FROM staging_{table_name}")
        rows = cursor.fetchall()
        print(type(rows))
        cursor.close()
        return {str(r[id_col]): r["checksum"] for r in rows if r[id_col] is not None}
    except Exception:
        cursor.close()
        return {}


def compute_change_flag(conn, table_name, id_col, id_value, incoming_checksum):
    base_record = get_active_base_record(conn, table_name, id_col, id_value)
    if base_record is None:
        return "I"
    elif base_record.get("checksum") != incoming_checksum:
        return "U"
    else:
        return "N"


def insert_staging_rows(conn, table_name, rows):
    if not rows:
        return
    cursor = conn.cursor()
    cols         = list(rows[0].keys())
    col_names    = ", ".join([f"`{c}`" for c in cols])
    placeholders = ", ".join(["%s"] * len(cols))
    sql          = f"INSERT INTO staging_{table_name} ({col_names}) VALUES ({placeholders})"
    values       = [tuple(row[c] for c in cols) for row in rows]
    cursor.executemany(sql, values)
    conn.commit()
    cursor.close()


def get_date_from_filename(fname):
    match = re.search(r"_(\d{6})\.", fname)
    return match.group(1) if match else None


def apply_transformations(row, table_name):
    row = standardize_column_casing(row)
    row = rename_columns(row)           # CHANGE 1: rename non-standard column names
    clean_whitespace(row)
    handle_nulls(row)
    capitalize_name_columns(row)        # CHANGE 2: now catches _nm columns too
    standardize_dates(row)
    validate_date_range(row, table_name)
    validate_email(row)
    validate_phone(row)
    standardize_gender(row)
    if table_name == "assessment":
        validate_assessment_score(row)
    return row


def standardize_column_casing(row):
    return {k.strip().lower().replace(" ", "_"): v for k, v in row.items()}


def rename_columns(row):
    renamed = {}
    for k, v in row.items():
        renamed[COLUMN_RENAME_MAP.get(k, k)] = v
    return renamed


def clean_whitespace(row):
    for k, v in row.items():
        if isinstance(v, str):
            row[k] = v.strip()


def handle_nulls(row):
    for k, v in row.items():
        if v is None or (isinstance(v, str) and v.strip().lower() in ("", "nan", "null", "none")):
            row[k] = None


def capitalize_name_columns(row):
    for k in row:
        if ("name" in k.lower() or k.lower().endswith("_nm")) and isinstance(row[k], str):
            row[k] = row[k].title()


def standardize_dates(row):
    for k in row:
        if ("date" in k.lower() or "dob" in k.lower()) and row[k]:
            try:
                dt = parser.parse(str(row[k]), dayfirst=True)
                row[k] = dt.strftime("%Y-%m-%d")
            except Exception:
                row[k] = None


def validate_date_range(row, table_name):
    pair = DATE_PAIR_MAP.get(table_name)
    if not pair:
        return
    start_col, end_col = pair
    start_val = row.get(start_col)
    end_val   = row.get(end_col)
    if start_val and end_val:
        try:
            if start_val > end_val:
                print(f"    invalid date range: {start_col}={start_val} > {end_col}={end_val} — nulled")
                row[start_col] = None
                row[end_col]   = None
        except Exception:
            pass


def validate_email(row):
    if "email" in row and row["email"]:
        row["email"] = str(row["email"]).strip().lower()
        pattern = r"^[\w\.-]+@[\w\.-]+\.\w{2,}$"
        if not re.match(pattern, row["email"]):
            row["email"] = None


def validate_phone(row):
    for k in row:
        if ("phone" in k.lower() or "mobile" in k.lower()) and row[k]:
            phone = re.sub(r"[\s\-]", "", str(row[k]).strip())
            if re.match(r"^\+?\d{7,15}$", phone):
                row[k] = phone
            else:
                row[k] = None


def standardize_gender(row):
    if "gender" in row and row["gender"] is not None:
        g = str(row["gender"]).strip().upper()
        gender_map = {"M": "M", "MALE": "M", "F": "F", "FEMALE": "F"}
        row["gender"] = gender_map.get(g, None)



def validate_assessment_score(row):
    if "score_prct" in row and row["score_prct"] is not None:
        try:
            score = float(row["score_prct"])
            row["score_prct"] = score if 0 < score < 100 else None
        except (ValueError, TypeError):
            row["score_prct"] = None



def run_staging():
    conn       = mysql.connector.connect(**DB_CONFIG)
    create_audit_table_2(conn)   # CHANGE (audit_table_2)

    load_key   = get_latest_load_key(conn)
    staging_ts = datetime.now()

    print(f"staging for load_key: {load_key}")
    print(f"staging timestamp: {staging_ts}\n")

    for table_name, (landing_table, id_col, archive_prefix) in TABLE_CONFIG.items():
        print(f"processing {table_name}...")
        staging_table_name = f"staging_{table_name}"   # CHANGE (audit_table_2)

        try:
            land_row = fetch_landing_data(conn, landing_table, load_key)
            if not land_row:
                print(f"  SKIPPP!! no data found in {landing_table}")
                # CHANGE (audit_table_2)
                write_to_audit_2(conn, staging_table_name, SCRIPT_NAME, staging_ts,
                                  load_key, "no_data", None, 0)
                continue

            transformed_rows = []
            for row in land_row:
                row.pop("landing_id", None)

                row = apply_transformations(row, table_name)

                if not mandatory_fields_present(row, table_name):
                    print(f"Skipping record - mandatory field missing: {row}")
                    continue

                transformed_rows.append(row)

            if not transformed_rows:
                print(f"No valid records for {table_name}")
                # CHANGE (audit_table_2)
                write_to_audit_2(conn, staging_table_name, SCRIPT_NAME, staging_ts,
                                  load_key, "no_valid_records", None, 0)
                continue

            data_cols = list(transformed_rows[0].keys())
            create_staging_table(conn, table_name, data_cols)

            existing_checksums = get_existing_staging_checksums(conn, table_name, id_col)

            staging_rows = []
            counts = {"inserted": 0, "updated": 0, "skipped_no_change": 0, "skipped_duplicate": 0}

            seen_ids = set()
            for curr_row in transformed_rows:

                id_value = curr_row.get(id_col)

                if id_value in seen_ids:
                    print(f"Skipping duplicate business key: {id_col} = {id_value}")
                    counts["skipped_duplicate"] += 1
                    continue

                seen_ids.add(id_value)
                
                checksum = compute_checksum(curr_row)
                id_value = curr_row.get(id_col, "UNKNOWN")

                if str(id_value) in existing_checksums and existing_checksums[str(id_value)] == checksum:
                    counts["skipped_duplicate"] += 1
                    continue

                change_flag = compute_change_flag(conn, table_name, id_col, id_value, checksum)

                if change_flag == "N":
                    counts["skipped_no_change"] += 1
                    continue

                s_key        = make_staging_key(id_value, staging_ts)
                key_col_name = f"{table_name}_key"
                final_row    = {key_col_name: s_key}
                final_row.update(curr_row)
                final_row["checksum"]    = checksum
                final_row["change_flag"] = change_flag

                staging_rows.append(final_row)
                if change_flag == "I":
                    counts["inserted"] += 1
                else:
                    counts["updated"] += 1

            insert_staging_rows(conn, table_name, staging_rows)
            print(
                f"  flagged I (new): {counts['inserted']}  "
                f"flagged U (changed): {counts['updated']}  "
                f"skipped N (no change): {counts['skipped_no_change']}  "
                f"skipped (duplicate in staging): {counts['skipped_duplicate']}"
            )

            # CHANGE (audit_table_2): row_count = rows actually written to staging
            write_to_audit_2(
                conn, staging_table_name, SCRIPT_NAME, staging_ts, load_key,
                "success", None, counts["inserted"] + counts["updated"]
            )

        except Exception as e:
            print(f"  error processing {table_name}: {e}")
            # CHANGE (audit_table_2)
            write_to_audit_2(conn, staging_table_name, SCRIPT_NAME, staging_ts,
                              load_key, "failed", None, 0)

    conn.close()
    print("\nstaging step done")


if __name__ == "__main__":
    run_staging()


# based on the latest load key in load control
# it fetches landing table records and applies transformations
# rename_columns handles non-standard column name variants (e.g. employee_id -> emp_id)
# capitalize_name_columns now catches both 'name' and '_nm' column suffixes
# checksum is computed from the transformed row for consistency across runs
# duplicate check: skip if same id+checksum already in staging
# change_flag check: compare checksum against active base record
#   I = new record, U = changed record, N = no change (skipped, not staged)
# only I and U rows are inserted into staging
# CHANGE (audit_table_2): every table processed here (success, no data,
# no valid records, or failed) now writes one row to audit_table_2 so the
# staging layer shows up in the same cross-layer audit trail as landing/base